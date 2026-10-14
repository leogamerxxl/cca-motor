"""Registrul Excel CCA_Etapa2_Calcule.xlsx: toate calculele P2 / P3 ca formule.

Apelat din calcule_p2_p3.py. Stilul urmează registrul din etapa 1:
text albastru = valoare de intrare, text negru = formulă.
"""
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_REF = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

DS = "'Date documentare'"
P2S = "'P2 Putere'"
P3S = "'P3 Energie'"


def antet(ws, rand, titluri, latimi=None):
    for c, h in enumerate(titluri, 1):
        cell = ws.cell(row=rand, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    if latimi:
        for c, w in enumerate(latimi, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def titlu(ws, text, nota=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if nota:
        ws["A2"] = nota
        ws["A2"].font = F_NOTE


def nume_definit(wb, nume, foaie, celula):
    ref = f"'{foaie}'!${re.sub(r'(\d+)', r'$\1', celula)}"
    wb.defined_names[nume] = DefinedName(nume, attr_text=ref)


# ---------------------------------------------------------------- date documentare
COL_INTRARE = [
    ("vehicul", "Vehicul", 34), ("P_max_kW", "P max [kW]", 9), ("v_max_kmh", "v max [km/h]", 9),
    ("t_acc_s", "t accel. [s]", 8), ("v_acc_kmh", "până la v [km/h]", 9), ("masa_kg", "Masă [kg]", 9),
    ("cuplu_Nm", "Cuplu [Nm]", 9), ("baterie_kWh", "Baterie [kWh]", 9),
    ("autonomie_min_km", "Autonomie min [km]", 10), ("autonomie_max_km", "Autonomie max [km]", 10),
    ("ciclu", "Ciclu", 8), ("consum_min_kWh100", "Consum min [kWh/100 km]", 10),
    ("consum_max_kWh100", "Consum max [kWh/100 km]", 10), ("Cx", "Cx", 7), ("A_m2", "S [m²]", 7),
    ("anv_spate_latime_mm", "Anvelopă spate lățime [mm]", 10), ("anv_spate_profil_pct", "Profil [%]", 7),
    ("janta_spate_inch", "Jantă [inch]", 7),
]
# coloane calculate: (cheie, antet, lățime, formulă cu {col:cheie} pentru rândul curent)
COL_CALC = [
    ("P_W", "P max [W]", 12, "={P_max_kW}*1000"),
    ("vmax_ms", "v max [m/s]", 9, "={v_max_kmh}/3.6"),
    ("a_ms2", "a medie [m/s²]", 9, "=({v_acc_kmh}/3.6)/{t_acc_s}"),
    ("bat_Ws", "Baterie [Ws]", 13, "={baterie_kWh}*3.6E6"),
    ("aut_km", "Autonomie medie [km]", 10, '=IF(COUNT({autonomie_min_km},{autonomie_max_km})<2,"",'
                                           'AVERAGE({autonomie_min_km},{autonomie_max_km}))'),
    ("cons_kWh100", "Consum mediu [kWh/100 km]", 10, '=IF(COUNT({consum_min_kWh100},{consum_max_kWh100})<2,"",'
                                                    'AVERAGE({consum_min_kWh100},{consum_max_kWh100}))'),
    ("cons_Wsm", "Consum [Ws/m]", 9, '=IF({cons_kWh100}="","",{cons_kWh100}/100*3600)'),
    ("r_static_m", "r roată (static) [m]", 9, '=IF(COUNT({anv_spate_latime_mm},{anv_spate_profil_pct},'
                                             '{janta_spate_inch})<3,"",({janta_spate_inch}*25.4/2+'
                                             '{anv_spate_latime_mm}*{anv_spate_profil_pct}/100)/1000)'),
]


def foaie_date(wb, tab):
    ws = wb.active
    ws.title = "Date documentare"
    titlu(ws, "Date din documentare (etapa 1) transformate în unități S.I.",
          "Se preiau datele medii din faza de documentare (suport P2/P3, slide 2). Celulele goale = valoare nedeclarată; "
          "media se face doar pe valorile existente. Rândul galben = vehiculul de referință.")
    ws["A3"] = "Text albastru = valoare preluată din sursă; text negru = formulă."
    ws["A3"].font = F_NOTE
    HR = 5
    toate = COL_INTRARE + [(k, h, w) for k, h, w, _ in COL_CALC]
    antet(ws, HR, ["Nr."] + [h for _, h, _ in toate], [5] + [w for _, _, w in toate])
    ws.row_dimensions[HR].height = 54
    col = {k: get_column_letter(i + 2) for i, (k, _, _) in enumerate(toate)}
    r0 = HR + 1
    rand_gt = None
    for i, v in enumerate(tab):
        r = r0 + i
        ws.cell(row=r, column=1, value=i + 1).font = F_BASE
        for k, _, _ in COL_INTRARE:
            val = v[k]
            if isinstance(val, float) and math.isnan(val):
                val = None
            c = ws[f"{col[k]}{r}"]
            c.value = val
            c.font = F_INPUT
        for k, _, _, fml in COL_CALC:
            c = ws[f"{col[k]}{r}"]
            c.value = re.sub(r"\{(\w+)\}", lambda m: f"{col[m.group(1)]}{r}", fml)
            c.font = F_BASE
        for cc in range(1, len(toate) + 2):
            ws.cell(row=r, column=cc).border = BORDER
            ws.cell(row=r, column=cc).alignment = WRAP
        if v["vehicul"].startswith("Mercedes-AMG GT 63"):
            rand_gt = r
            for cc in range(1, len(toate) + 2):
                ws.cell(row=r, column=cc).fill = FILL_REF
    r1 = r0 + len(tab) - 1
    rm, rn = r1 + 1, r1 + 2
    ws.cell(row=rm, column=2, value="MEDIE (vehicul mediu)").font = F_BOLD
    ws.cell(row=rn, column=2, value="Număr de valori").font = F_NOTE
    doar_wltp = ("autonomie_min_km", "autonomie_max_km", "consum_min_kWh100", "consum_max_kWh100",
                 "aut_km", "cons_kWh100", "cons_Wsm")
    Lc = col["ciclu"]
    for k, _, _ in toate:
        if k in ("vehicul", "ciclu"):
            continue
        L = col[k]
        if k in doar_wltp:
            ws[f"{L}{rm}"] = f'=AVERAGEIF({Lc}{r0}:{Lc}{r1},"WLTP",{L}{r0}:{L}{r1})'
            ws[f"{L}{rn}"] = f'=COUNTIFS({Lc}{r0}:{Lc}{r1},"WLTP",{L}{r0}:{L}{r1},">=0")'
        else:
            ws[f"{L}{rm}"] = f"=AVERAGE({L}{r0}:{L}{r1})"
            ws[f"{L}{rn}"] = f"=COUNT({L}{r0}:{L}{r1})"
        ws[f"{L}{rm}"].font = F_BOLD
        ws[f"{L}{rm}"].fill = FILL_KEY
        ws[f"{L}{rn}"].font = F_NOTE
        for rr in (rm, rn):
            ws[f"{L}{rr}"].border = BORDER
    for k in ("P_W", "bat_Ws"):
        for rr in range(r0, rm + 1):
            ws[f"{col[k]}{rr}"].number_format = "0.000E+00"
    for k in ("a_ms2", "vmax_ms", "cons_Wsm", "aut_km", "cons_kWh100"):
        for rr in range(r0, rm + 1):
            ws[f"{col[k]}{rr}"].number_format = "0.00"
    for k in ("r_static_m", "Cx"):
        for rr in range(r0, rm + 1):
            ws[f"{col[k]}{rr}"].number_format = "0.000"
    ws.freeze_panes = ws.cell(row=r0, column=3)
    nota = rn + 2
    ws.cell(row=nota, column=2, value=(
        "Note: Lucid Air Sapphire – 1,89 s este 0–60 mph (96,56 km/h). Autonomia și consumul se mediază doar pe "
        "valorile WLTP (Xiaomi și Yangwang declară ciclul chinezesc CLTC, mai optimist). Lotus nu publică un consum "
        "WLTP separat pentru Emeya R. Tesla și Maserati: autonomie și consum din EV Database (sursă secundară). "
        "Raza roții se calculează din anvelopa punții spate: r = D_jantă/2 + lățime·profil.")).font = F_NOTE
    return col, rm, rand_gt


# ---------------------------------------------------------------- ipoteze
IPOTEZE = [
    ("c_g", "Accelerația gravitațională g", "m/s²", "g", "Valoare uzuală (g_n = 9,80665 m/s²)"),
    ("c_rho", "Densitatea aerului ρ (condiții normale: 0 °C, 101 325 Pa)", "kg/m³", "rho",
     "ρ = p/(R_s·T) = 101325/(287,05·273,15)"),
    ("c_mu_f", "Coeficient de frecare la rulare µf", "-", "mu_f", "Exemplul din curs (P2, slide 5)"),
    ("c_mu_a", "Coeficient de aderență µa", "-", "mu_a", "Exemplul din curs (P2, slide 5)"),
    ("c_alpha", "Unghiul de înclinare al drumului α", "rad", "alpha", "Drum drept (simplificare curs)"),
    ("c_Vv", "Viteza vântului Vv", "m/s", "Vv", "Fără vânt (simplificare curs)"),
    ("c_r", "Raza dinamică a roții r", "m", "r", "Media razelor anvelopelor spate ≈ 0,36 m × 0,97 ≈ 0,35 m"),
    ("c_pp", "Pierderi în reductor Pp / Pax", "-", "pp", "Pp = Pax/100 (curs, P2 slide 7)"),
    ("c_Va_kmh", "Viteza Va (sfârșitul accelerării)", "km/h", "Va_kmh", "0–100 km/h"),
    ("c_n_max", "Turația maximă a mașinii electrice la v_max", "rot/min", "n_max",
     "Alegere de proiectare (motoarele clasei: 13 000–20 000 rot/min)"),
    ("c_V1_kmh", "Viteza constantă a ciclului P3, V1", "km/h", "V1_kmh", "Curs, P3 slide 4"),
    ("c_eta", "Randament η_el-mec = η_mec-el", "-", "eta", "Curs, P3 slide 8"),
    ("c_rez", "Rezervă de energie în baterie", "-", "rezerva", "Curs, P3 slide 8"),
    ("c_U_tinta", "Tensiunea nominală țintă a bateriei", "V", "U_bat_tinta", "Arhitectura de 800 V a clasei"),
    ("c_k_masa", "Masa maximă a bateriei / masa vehiculului", "-", "k_masa_bat", "Curs, P3 slide 9"),
]
CELULA_CAMPURI = [
    ("cel_U_nom", "Tensiune nominală celulă", "V", "U_nom"),
    ("cel_U_max", "Tensiune maximă (încărcare)", "V", "U_max"),
    ("cel_U_min", "Tensiune minimă (descărcare)", "V", "U_min"),
    ("cel_Cap", "Capacitate nominală", "Ah", "Cap_Ah"),
    ("cel_m", "Masă celulă", "kg", "m_kg"),
    ("cel_I", "Curent maxim de descărcare continuă", "A", "I_max_cont"),
    ("cel_kcp", "Raport masă celule / masă pachet", "-", "k_celula_pachet"),
]


def foaie_ipoteze(wb, C, cel):
    ws = wb.create_sheet("Ipoteze")
    titlu(ws, "Ipoteze și constante de calcul",
          "Celulele albastre se pot modifica; toate foile de calcul folosesc numele definite din coloana A.")
    antet(ws, 4, ["Nume definit", "Mărime", "Valoare", "Unitate", "Justificare / sursă"], [14, 52, 12, 10, 70])
    r = 5
    for nume, et, um, cheie, just in IPOTEZE:
        ws.cell(row=r, column=1, value=nume).font = F_NOTE
        ws.cell(row=r, column=2, value=et).font = F_BASE
        ws.cell(row=r, column=3, value=C[cheie]).font = F_INPUT
        ws.cell(row=r, column=4, value=um).font = F_BASE
        ws.cell(row=r, column=5, value=just).font = F_BASE
        for cc in range(1, 6):
            ws.cell(row=r, column=cc).border = BORDER
        nume_definit(wb, nume, "Ipoteze", f"C{r}")
        r += 1
    r += 1
    ws.cell(row=r, column=2, value=f"Celula Li-ion aleasă: {cel['nume']}").font = F_BOLD
    ws.cell(row=r, column=5, value=cel["sursa"]).font = F_LINK
    ws.cell(row=r, column=5).hyperlink = cel["sursa"]
    r += 1
    for nume, et, um, cheie in CELULA_CAMPURI:
        ws.cell(row=r, column=1, value=nume).font = F_NOTE
        ws.cell(row=r, column=2, value=et).font = F_BASE
        ws.cell(row=r, column=3, value=cel[cheie]).font = F_INPUT
        ws.cell(row=r, column=4, value=um).font = F_BASE
        ws.cell(row=r, column=5, value=cel.get("note_" + cheie, "")).font = F_BASE
        for cc in range(1, 6):
            ws.cell(row=r, column=cc).border = BORDER
        nume_definit(wb, nume, "Ipoteze", f"C{r}")
        r += 1


# ---------------------------------------------------------------- foi de calcul cu rânduri (mediu | GT 63)
class FoaieCalcul:
    """Foaie cu un rând pe mărime și o coloană pe set de date (E = vehicul mediu, F = GT 63)."""

    COLS = ("E", "F")

    def __init__(self, wb, nume, titlu_, nota):
        self.ws = wb.create_sheet(nume)
        self.nume = nume
        titlu(self.ws, titlu_, nota)
        antet(self.ws, 4, ["Pas", "Mărime", "Simbol", "Unitate", "Vehicul mediu", "GT 63 (referință)",
                           "Formulă / observație"], [6, 46, 13, 9, 16, 16, 70])
        self.r = 5
        self.rand = {}

    def sectiune(self, text):
        c = self.ws.cell(row=self.r, column=1, value=text)
        c.font = F_BOLD
        self.r += 1

    def rand_(self, cheie, pas, et, simbol, um, formula, obs="", fmt="0.000", cheie_rez=False, intrare=False):
        ws, r = self.ws, self.r
        self.rand[cheie] = r
        for cc, val in ((1, pas), (2, et), (3, simbol), (4, um), (7, obs)):
            c = ws.cell(row=r, column=cc, value=val)
            c.font = F_BASE if cc != 7 else F_NOTE
        for i, L in enumerate(self.COLS):
            fml = formula[i] if isinstance(formula, (list, tuple)) else formula
            fml = re.sub(r"\{(\w+)\}", lambda m: f"{L}{self.rand[m.group(1)]}", fml)
            c = ws[f"{L}{r}"]
            c.value = fml
            c.font = F_INPUT if intrare else (F_BOLD if cheie_rez else F_BASE)
            c.number_format = fmt
            if cheie_rez:
                c.fill = FILL_KEY
        for cc in range(1, 8):
            ws.cell(row=r, column=cc).border = BORDER
            ws.cell(row=r, column=cc).alignment = WRAP
        self.r += 1
        return r

    def ref(self, cheie, i):
        return f"'{self.nume}'!{self.COLS[i]}{self.rand[cheie]}"


def foaie_p2(wb, col, rm, rgt):
    F = FoaieCalcul(wb, "P2 Putere", "P2 – Calculul puterii mașinii electrice (pașii 1–9)",
                    "Strategie: se pleacă de la forțele rezistente pe vehicul: Ft → Mr → Pr → Pax. "
                    "Simplificări: α = 0, Vv = 0, deplasare în linie dreaptă (Vr = V).")
    src = lambda k: [f"={DS}!{col[k]}{rm}", f"={DS}!{col[k]}{rgt}"]  # noqa: E731
    F.sectiune("Date de intrare (din documentare, S.I.)")
    F.rand_("M", "", "Masa vehiculului", "M", "kg", src("masa_kg"), "medie / rândul GT 63", "0.0")
    F.rand_("Cx", "", "Coeficient aerodinamic", "Cx", "-", src("Cx"), "", "0.0000")
    F.rand_("S", "", "Suprafața frontală", "S", "m²", src("A_m2"), "", "0.000")
    F.rand_("vmax", "", "Viteza maximă", "Vmax", "m/s", src("vmax_ms"), "", "0.000")
    F.rand_("a_doc", "", "Accelerația medie 0–100 km/h (documentare)", "a_doc", "m/s²", src("a_ms2"), "", "0.000")
    F.rand_("P_doc", "", "Puterea maximă declarată", "P_doc", "W", src("P_W"), "", "0")
    F.rand_("r", "", "Raza roții", "r", "m", "=c_r", "Ipoteze", "0.000")
    F.sectiune("Pașii 1–4: forța de tracțiune în punctul (Va, ta) și aderența")
    F.rand_("Va", "1–2", "Viteza la sfârșitul accelerării", "Va", "m/s", "=c_Va_kmh/3.6", "Va = 100 km/h")
    F.rand_("ta_doc", "2", "Timpul de accelerare (documentare)", "ta", "s", "={Va}/{a_doc}", "ta = Va / a")
    F.rand_("Ff", "3", "Forța de frecare la rulare", "Ff", "N", "=c_mu_f*{M}*c_g*COS(c_alpha)", "Ff = µf·M·g·cos α", "0.00")
    F.rand_("Fra_Va", "3", "Rezistența aerodinamică la Va", "Fra", "N", "=0.5*c_rho*{S}*{Cx}*({Va}-c_Vv)^2",
            "Fra = ½·ρ·S·Cx·(V − Vv)²", "0.00")
    F.rand_("Gt", "3", "Componenta tangențială a greutății", "Gt", "N", "={M}*c_g*SIN(c_alpha)", "Gt = M·g·sin α", "0.00")
    F.rand_("Fi_doc", "3", "Forța de inerție cu a_doc", "M·a", "N", "={M}*{a_doc}", "", "0.00")
    F.rand_("Ft_doc", "3", "Forța de tracțiune necesară", "Ft", "N", "={Fra_Va}+{Ff}+{Gt}+{Fi_doc}",
            "Ft = Fra + Ff + Gt + M·a", "0.00")
    F.rand_("Fa", "4", "Forța de aderență", "Fa", "N", "=c_mu_a*{M}*c_g*COS(c_alpha)", "Fa = µa·M·g·cos α", "0.00")
    F.rand_("ader", "4", "Aderența respectată? (Ft ≤ Fa)", "", "", '=IF({Ft_doc}<={Fa},"DA","NU")',
            "Dacă NU: Ft = Fa și se recalculează accelerația", "@")
    F.rand_("a", "4", "Accelerația realizabilă", "a", "m/s²",
            "=IF({Ft_doc}<={Fa},{a_doc},({Fa}-{Fra_Va}-{Ff}-{Gt})/{M})", "a = (Fa − Fra − Ff − Gt)/M", "0.000", True)
    F.rand_("Ft", "4", "Forța de tracțiune de calcul", "Ft", "N", "=IF({Ft_doc}<={Fa},{Ft_doc},{Fa})", "", "0.00")
    F.rand_("ta", "4", "Timpul 0–100 km/h realizabil", "ta", "s", "={Va}/{a}", "", "0.000")
    F.sectiune("Pasul 5: cuplul și puterea la roată; puterea la axul mașinii")
    F.rand_("Om_r_Va", "5", "Viteza unghiulară a roții la Va", "Ωr", "rad/s", "={Va}/{r}", "Ωr = Va / r", "0.000")
    F.rand_("Mr", "5", "Cuplul la roată (total)", "Mr", "Nm", "={r}*{Ft}", "Mr = r·Ft", "0.0")
    F.rand_("Pr", "5", "Puterea la roată", "Pr", "W", "={Mr}*{Om_r_Va}", "Pr = Mr·Ωr = Ft·Vr", "0")
    F.rand_("Pax", "5", "Puterea la axul mașinii electrice", "Pax", "W", "={Pr}/(1-c_pp)",
            "Pax = Pr + Pp, Pp = Pax/100  ⇒  Pax = Pr/0,99", "0", True)
    F.rand_("Pp", "5", "Pierderi în reductor și transmisie", "Pp", "W", "={Pax}-{Pr}", "", "0")
    F.sectiune("Pasul 6: reductorul și cuplul la axul mașinii")
    F.rand_("Om_r_max", "6", "Viteza unghiulară a roții la Vmax", "Ωr,max", "rad/s", "={vmax}/{r}", "", "0.000")
    F.rand_("Om_ax_max", "6", "Viteza unghiulară maximă a mașinii", "Ωax,max", "rad/s", "=c_n_max*2*PI()/60",
            "din n_max ales (Ipoteze)", "0.000")
    F.rand_("i_rd", "6", "Raportul reductorului", "i_rd", "-", "={Om_r_max}/{Om_ax_max}", "i_rd = Ωr / Ωax", "0.00000",
            True)
    F.rand_("k_rd", "6", "Raport de reducere (n_motor / n_roată)", "1/i_rd", "-", "=1/{i_rd}", "", "0.000")
    F.rand_("Om_ax_Va", "6", "Viteza unghiulară a mașinii la Va (de bază)", "Ωax", "rad/s", "={Om_r_Va}/{i_rd}", "",
            "0.000")
    F.rand_("n_ax_Va", "6", "Turația de bază a mașinii", "n_b", "rot/min", "={Om_ax_Va}*60/(2*PI())", "", "0")
    F.rand_("Max", "6", "Cuplul maxim la axul mașinii", "Max", "Nm", "={Pax}/{Om_ax_Va}", "Max = Pax / Ωax", "0.0", True)
    F.sectiune("Pașii 7–9: rulare cu Vmax constantă (a = 0) – puterea nominală")
    F.rand_("Fra_vmax", "7", "Rezistența aerodinamică la Vmax", "Fra", "N", "=0.5*c_rho*{S}*{Cx}*({vmax}-c_Vv)^2", "",
            "0.00")
    F.rand_("Ft_vmax0", "7", "Forța de tracțiune la Vmax", "Ft", "N", "={Fra_vmax}+{Ff}+{Gt}", "Ft = Fra + Ff + Gt",
            "0.00")
    F.rand_("ader_vmax", "8", "Aderența respectată la Vmax?", "", "", '=IF({Ft_vmax0}<={Fa},"DA","NU")',
            "Dacă NU se reduce Vmax până la Ft = Fa", "@")
    F.rand_("vmax_ef", "8", "Viteza maximă de calcul", "Vmax", "m/s",
            "=IF({Ft_vmax0}<={Fa},{vmax},SQRT(({Fa}-{Ff}-{Gt})/(0.5*c_rho*{S}*{Cx}))+c_Vv)", "", "0.000")
    F.rand_("Ft_vmax", "8", "Forța de tracțiune de calcul la Vmax", "Ft", "N", "=IF({Ft_vmax0}<={Fa},{Ft_vmax0},{Fa})",
            "", "0.00")
    F.rand_("Mr_vmax", "8", "Cuplul la roată la Vmax", "Mr", "Nm", "={r}*{Ft_vmax}", "", "0.0")
    F.rand_("Pr_vmax", "8", "Puterea la roată la Vmax", "Pr", "W", "={Ft_vmax}*{vmax_ef}", "", "0")
    F.rand_("Pax_nom", "8", "Puterea la axul mașinii la Vmax", "Pax", "W", "={Pr_vmax}/(1-c_pp)", "", "0", True)
    F.rand_("Max_vmax", "8", "Cuplul la axul mașinii la Vmax", "Max", "Nm", "={Pax_nom}/{Om_ax_max}", "", "0.0")
    F.rand_("Pmax_kW", "9", "PUTEREA MAXIMĂ a mașinii electrice", "P_max", "kW", "={Pax}/1000", "din pasul 5", "0.0",
            True)
    F.rand_("Pnom_kW", "9", "PUTEREA NOMINALĂ a mașinii electrice", "P_nom", "kW", "={Pax_nom}/1000", "din pasul 8",
            "0.0", True)
    F.sectiune("Verificarea coerenței cu documentarea")
    F.rand_("Pax_fara", "", "Pax cu accelerația documentată (fără limita de aderență)", "", "kW",
            "={Ft_doc}*{Va}/(1-c_pp)/1000", "aceeași formulă cu a_doc", "0.0")
    F.rand_("Pdoc_kW", "", "Puterea maximă declarată (documentare)", "", "kW", "={P_doc}/1000", "", "0.0")
    F.rand_("raport", "", "Vmax / Va (domeniu de putere constantă necesar)", "", "-", "={vmax_ef}/{Va}", "", "0.00")
    return F


def foaie_p3(wb, col, rm, rgt, F2):
    F = FoaieCalcul(wb, "P3 Energie", "P3 – Necesarul de energie electrică (pașii 1–14)",
                    "Ciclu: accelerare cu a până la V1 = 50 km/h, rulare cu V1 constantă, decelerare (df = d1). "
                    "Distanța totală a ciclului = autonomia medie din documentare.")
    src = lambda k: [f"={DS}!{col[k]}{rm}", f"={DS}!{col[k]}{rgt}"]  # noqa: E731
    F.sectiune("Date de intrare")
    F.rand_("M", "", "Masa vehiculului", "M", "kg", src("masa_kg"), "", "0.0")
    F.rand_("Cx", "", "Coeficient aerodinamic", "Cx", "-", src("Cx"), "", "0.0000")
    F.rand_("S", "", "Suprafața frontală", "S", "m²", src("A_m2"), "", "0.000")
    F.rand_("aut", "", "Cerința de autonomie", "d", "m", [f"={DS}!{col['aut_km']}{rm}*1000",
                                                           f"={DS}!{col['aut_km']}{rgt}*1000"], "media din documentare",
            "0")
    F.rand_("a_P2", "2", "Accelerația calculată în P2", "a", "m/s²", [f"={F2.ref('a', 0)}", f"={F2.ref('a', 1)}"],
            "sesiunea precedentă (P2, pasul 4)", "0.000")
    F.sectiune("Pașii 1–6: puterea P1 (cu a), P1' (fără a), timpul t1")
    F.rand_("V1", "1", "Viteza constantă a ciclului", "V1", "m/s", "=c_V1_kmh/3.6", "V1 = 50 km/h", "0.0000")
    F.rand_("Ff", "3", "Forța de frecare la rulare", "Ff", "N", "=c_mu_f*{M}*c_g*COS(c_alpha)", "", "0.000")
    F.rand_("Fra", "3", "Rezistența aerodinamică la V1", "Fra", "N", "=0.5*c_rho*{S}*{Cx}*({V1}-c_Vv)^2", "", "0.000")
    F.rand_("Gt", "3", "Componenta tangențială a greutății", "Gt", "N", "={M}*c_g*SIN(c_alpha)", "", "0.000")
    F.rand_("Fa", "4", "Forța de aderență", "Fa", "N", "=c_mu_a*{M}*c_g*COS(c_alpha)", "", "0.00")
    F.rand_("Ft1_0", "3", "Ft1 cu accelerația din P2", "Ft1", "N", "={Fra}+{Ff}+{Gt}+{M}*{a_P2}", "", "0.00")
    F.rand_("ader", "4", "Aderența respectată? (Ft1 ≤ Fa)", "", "", '=IF({Ft1_0}<={Fa},"DA","NU")', "", "@")
    F.rand_("a", "4", "Accelerația folosită", "a", "m/s²", "=IF({Ft1_0}<={Fa},{a_P2},({Fa}-{Fra}-{Ff}-{Gt})/{M})",
            "", "0.000")
    F.rand_("Ft1", "3", "Forța de tracțiune în (V1, t1)", "Ft1", "N", "={Fra}+{Ff}+{Gt}+{M}*{a}",
            "Ft1 = Fra + Ff + Gt + M·a", "0.00")
    F.rand_("P1", "3", "Puterea P1 (cu accelerație)", "P1", "W", "={Ft1}*{V1}", "P1 = Ft1·V1", "0.0", True)
    F.rand_("Ft1p", "5", "Forța de tracțiune la V1 constant", "Ft1'", "N", "={Fra}+{Ff}+{Gt}", "Ft1' = Fra + Ff + Gt",
            "0.000")
    F.rand_("P1p", "5", "Puterea P1' (fără accelerație)", "P1'", "W", "={Ft1p}*{V1}", "P1' = Ft1'·V1", "0.00", True)
    F.rand_("t1", "6", "Timpul de accelerare până la V1", "t1", "s", "={V1}/{a}", "t1 = V1 / a", "0.00000", True)
    F.sectiune("Pașii 7–8: puterea medie și energia pe [0, t1]")
    F.rand_("Pmed", "7", "Puterea medie (aproximare)", "Pmed", "W", "=(0+{P1})/2", "Pmed = (P0 + P1)/2, P0 = 0", "0.0")
    for n in (100, 1000):
        sh = f"'P3 n={n}'"
        F.rand_(f"P01_{n}", "7", f"Puterea medie discretizată, n = {n}", "P01-med", "W",
                [f"={sh}!C4", f"={sh}!H4"], f"(1/n)·Σ Fti·Vti (foaia P3 n={n})", "0.0")
        F.rand_(f"W01_{n}", "8", f"Energia W01 discretizată, n = {n}", "W01", "Ws",
                [f"={sh}!C5", f"={sh}!H5"], "Σ Pti·Δti", "0.0")
    F.rand_("W01_ex", "8", "Energia W01 exactă (integrală analitică)", "W01", "Ws",
            "=({Ff}+{Gt}+{M}*{a})*{a}*{t1}^2/2+0.5*c_rho*{S}*{Cx}*{a}^3*{t1}^4/4",
            "∫(A + B·v²)·v dt, v = a·t (Vv = 0)", "0.0")
    F.rand_("err100", "8", "Eroarea W01 pentru n = 100", "", "%", "=({W01_100}-{W01_ex})/{W01_ex}*100", "", "0.0000")
    F.rand_("err1000", "8", "Eroarea W01 pentru n = 1000", "", "%", "=({W01_1000}-{W01_ex})/{W01_ex}*100", "", "0.0000")
    F.rand_("Wmed", "8", "Energia cu puterea medie", "Wmed", "Ws", "={Pmed}*{t1}", "Wmed = Pmed·t1", "0.0")
    F.rand_("W01med", "8", "Energia cu P01-med (n = 1000)", "W01-med", "Ws", "={P01_1000}*{t1}", "W01-med = P01-med·t1",
            "0.0")
    F.rand_("W01", "8", "Energia W01 folosită mai departe (n = 1000)", "W01", "Ws", "={W01_1000}", "", "0.0", True)
    F.sectiune("Pașii 9–12: distanțe și energia mecanică pe ciclu")
    F.rand_("d1", "9", "Distanța de accelerare", "d1", "m", "={V1}*{t1}/2", "d1 = V1·t1/2", "0.0000")
    F.rand_("df", "10", "Distanța de decelerare", "df", "m", "={d1}", "df = d1", "0.0000")
    F.rand_("d2", "10", "Distanța la viteză constantă", "d2", "m", "={aut}-{d1}-{df}", "autonomie = d1 + d2 + df", "0.000")
    F.rand_("t12", "10", "Durata rulării cu V1", "t2 − t1", "s", "={d2}/{V1}", "t2 − t1 = d2 / V1", "0.00")
    F.rand_("t2", "10", "Momentul t2", "t2", "s", "={t1}+{t12}", "", "0.00")
    F.rand_("tf", "10", "Durata totală a ciclului", "tf", "s", "={t2}+{t1}", "decelerare simetrică", "0.00")
    F.rand_("W12", "11", "Energia la viteză constantă", "W12", "Ws", "={P1p}*{t12}", "W12 = P1'·(t2 − t1)", "0.000E+00",
            True)
    F.rand_("W2f", "12", "Energia recuperabilă la decelerare", "W2f", "Ws", "={W01}", "W2f = W01", "0.0")
    F.rand_("Wmec", "12", "Energia mecanică pe ciclu", "Wmec", "Ws", "={W01}+{W12}-{W2f}", "Wmec = W01 + W12 − W2f",
            "0.000E+00")
    F.sectiune("Pașii 13–14: energia electrică în baterie")
    F.rand_("Wstoc", "13", "Energia electrică stocată necesară", "Wstocată", "Ws",
            "={W01}/c_eta+{W12}/c_eta-c_eta*{W2f}", "W01/η + W12/η − η·W2f", "0.000E+00", True)
    F.rand_("Wbat", "14", "Energia totală în baterie (rezervă 30 %)", "Wbat", "Ws", "={Wstoc}/(1-c_rez)",
            "Wbat = Wstocată + 0,3·Wbat", "0.000E+00", True)
    F.rand_("Wbat_kWh", "14", "Energia totală în baterie", "Wbat", "kWh", "={Wbat}/3600000", "", "0.000", True)
    F.rand_("Wstoc_kWh", "", "Energia stocată necesară", "", "kWh", "={Wstoc}/3600000", "", "0.000")
    F.rand_("cons", "", "Consum echivalent al ciclului", "", "kWh/100 km", "={Wstoc}/{aut}*100/3600", "", "0.00")
    F.rand_("cons_doc", "", "Consum mediu din documentare", "", "kWh/100 km",
            [f"={DS}!{col['cons_kWh100']}{rm}", f"={DS}!{col['cons_kWh100']}{rgt}"], "WLTP / EPA / CLTC", "0.00")
    F.rand_("bat_doc", "", "Capacitatea bateriei din documentare", "", "kWh",
            [f"={DS}!{col['baterie_kWh']}{rm}", f"={DS}!{col['baterie_kWh']}{rgt}"], "", "0.0")
    return F


def foaie_discretizare(wb, n, F3):
    ws = wb.create_sheet(f"P3 n={n}")
    titlu(ws, f"P3 – discretizarea intervalului [0, t1] în n = {n} intervale egale",
          "Pe fiecare interval [t(i−1), t(i)] forța Fti și viteza Vti sunt constante, egale cu valorile din t(i). "
          "Rândul i = 0 este t = 0 (P = 0), deci Σ pe cele n+1 puncte = Σ pe cele n intervale.")
    blocuri = (("A", 0, "Vehicul mediu"), ("F", 1, "GT 63"))
    r0 = 9
    for c0, k, et in blocuri:
        cA = c0
        ci = [get_column_letter(ord(cA) - 64 + j) for j in range(5)]  # i, t, V, F, P
        ws[f"{ci[0]}3"] = et
        ws[f"{ci[0]}3"].font = F_BOLD
        rot = {"t1": F3.ref("t1", k), "a": F3.ref("a", k), "M": F3.ref("M", k), "S": F3.ref("S", k),
               "Cx": F3.ref("Cx", k), "Ff": F3.ref("Ff", k), "Gt": F3.ref("Gt", k)}
        ws[f"{ci[1]}4"] = "P01-med [W]"
        ws[f"{ci[2]}4"] = f"=SUM({ci[4]}{r0 + 1}:{ci[4]}{r0 + n})/{n}"
        ws[f"{ci[1]}5"] = "W01 [Ws]"
        ws[f"{ci[2]}5"] = f"=SUM({ci[4]}{r0}:{ci[4]}{r0 + n})*{ci[2]}6"
        ws[f"{ci[1]}6"] = "Δt [s]"
        ws[f"{ci[2]}6"] = f"={rot['t1']}/{n}"
        ws[f"{ci[1]}7"] = "a [m/s²]"
        ws[f"{ci[2]}7"] = f"={rot['a']}"
        for rr in range(4, 8):
            ws[f"{ci[1]}{rr}"].font = F_BOLD
            ws[f"{ci[2]}{rr}"].font = F_BOLD
            ws[f"{ci[2]}{rr}"].number_format = "0.000000" if rr >= 6 else "0.000"
        for j, h in enumerate(["i", "t_i [s]", "V_ti [m/s]", "F_ti [N]", "P_ti [W]"]):
            c = ws[f"{ci[j]}{r0 - 1}"]
            c.value = h
            c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER
        dt, a = f"${ci[2]}$6", f"${ci[2]}$7"
        for i in range(n + 1):
            r = r0 + i
            ws[f"{ci[0]}{r}"] = i
            ws[f"{ci[1]}{r}"] = f"={ci[0]}{r}*{dt}"
            ws[f"{ci[2]}{r}"] = f"={a}*{ci[1]}{r}"
            ws[f"{ci[3]}{r}"] = (f"=0.5*c_rho*{rot['S']}*{rot['Cx']}*({ci[2]}{r}-c_Vv)^2+{rot['Ff']}+{rot['Gt']}"
                                 f"+{rot['M']}*{a}")
            ws[f"{ci[4]}{r}"] = f"={ci[3]}{r}*{ci[2]}{r}"
            for j in range(5):
                ws[f"{ci[j]}{r}"].font = F_BASE
            ws[f"{ci[1]}{r}"].number_format = "0.000000"
            ws[f"{ci[2]}{r}"].number_format = "0.0000"
            ws[f"{ci[3]}{r}"].number_format = "0.00"
            ws[f"{ci[4]}{r}"].number_format = "0.0"
        for j, w in enumerate([7, 12, 12, 12, 13]):
            ws.column_dimensions[ci[j]].width = w
    ws.freeze_panes = ws[f"A{r0}"]


def foaie_baterie(wb, F2, F3, col, rm, rgt):
    F = FoaieCalcul(wb, "Baterie", "P3 – Calculul bateriei electrice (celule Li-ion)",
                    "W_celulă = Cap·U; N = Wbat / W_celulă. Serie → tensiunea bateriei; paralel → curentul și capacitatea. "
                    "Limitare: masa bateriei ≤ 30 % din masa vehiculului.")
    F.sectiune("Date de intrare")
    F.rand_("Wbat", "", "Energia necesară în baterie (P3, pasul 14)", "Wbat", "Ws",
            [f"={F3.ref('Wbat', 0)}", f"={F3.ref('Wbat', 1)}"], "", "0.000E+00")
    F.rand_("Pmax", "", "Puterea maximă la axul mașinii (P2, pasul 5)", "P_max", "W",
            [f"={F2.ref('Pax', 0)}", f"={F2.ref('Pax', 1)}"], "", "0")
    F.rand_("M", "", "Masa vehiculului", "M", "kg", [f"={DS}!{col['masa_kg']}{rm}", f"={DS}!{col['masa_kg']}{rgt}"], "",
            "0.0")
    F.sectiune("Numărul de celule")
    F.rand_("Wcel", "", "Energia unei celule", "W_celulă", "Ws", "=cel_Cap*3600*cel_U_nom", "Cap [Ah→As] × U [V]", "0")
    F.rand_("Nmin", "", "Numărul minim de celule (energie)", "N_min", "-", "=ROUNDUP({Wbat}/{Wcel},0)", "", "0")
    F.rand_("Ns", "", "Celule în serie", "Ns", "-", "=ROUND(c_U_tinta/cel_U_nom,0)", "U_bat ≈ 800 V", "0", True)
    F.rand_("Unom", "", "Tensiunea nominală a bateriei", "U_bat", "V", "={Ns}*cel_U_nom", "", "0.0")
    F.rand_("Umax", "", "Tensiunea maximă (complet încărcată)", "", "V", "={Ns}*cel_U_max", "", "0.0")
    F.rand_("Umin", "", "Tensiunea minimă (descărcată)", "", "V", "={Ns}*cel_U_min", "", "0.0")
    F.rand_("Npe", "", "Ramuri în paralel – din energie", "Np", "-", "=ROUNDUP({Nmin}/{Ns},0)", "", "0")
    F.rand_("Pel", "", "Puterea electrică maximă cerută bateriei", "", "W", "={Pmax}/c_eta", "P_max / η", "0")
    F.rand_("Ibat", "", "Curentul maxim al bateriei", "I_bat", "A", "={Pel}/{Unom}", "", "0.0")
    F.rand_("Npc", "", "Ramuri în paralel – din curent", "Np", "-", "=ROUNDUP({Ibat}/cel_I,0)", "", "0")
    F.rand_("Np", "", "Ramuri în paralel (alese)", "Np", "-", "=MAX({Npe},{Npc})", "", "0", True)
    F.rand_("N", "", "Numărul total de celule", "N", "-", "={Ns}*{Np}", "", "0", True)
    F.rand_("Wreal", "", "Energia instalată", "", "kWh", "={N}*{Wcel}/3600000", "", "0.00", True)
    F.rand_("Cap", "", "Capacitatea bateriei", "", "Ah", "={Np}*cel_Cap", "", "0.0")
    F.rand_("Icel", "", "Curentul pe celulă la P_max", "", "A", "={Ibat}/{Np}", "", "0.00")
    F.sectiune("Verificarea masei")
    F.rand_("mcel", "", "Masa celulelor", "", "kg", "={N}*cel_m", "", "0.0", True)
    F.rand_("mpac", "", "Masa estimată a pachetului", "", "kg", "={mcel}/cel_kcp", "celule / raport celule-pachet", "0.0")
    F.rand_("mlim", "", "Limita: 30 % din masa vehiculului", "", "kg", "=c_k_masa*{M}", "", "0.0")
    F.rand_("pcel", "", "Masa celulelor / masa vehiculului", "", "%", "={mcel}/{M}*100", "", "0.0")
    F.rand_("ppac", "", "Masa pachetului / masa vehiculului", "", "%", "={mpac}/{M}*100", "", "0.0")
    F.rand_("ok", "", "Limita de 30 % respectată (pachet)?", "", "", '=IF({mpac}<={mlim},"DA","NU")', "", "@")
    return F


def foaie_sinteza(wb, F2, F3, FB):
    ws = wb.create_sheet("Sinteză", 1)
    titlu(ws, "Sinteză – date pentru proiectare (tema de proiectare)",
          "Valorile provin din foile P2, P3 și Baterie. Vehiculul mediu este baza temei; GT 63 este comparația.")
    antet(ws, 4, ["Subsistem", "Mărime", "Unitate", "Vehicul mediu", "GT 63 (referință)"], [26, 52, 12, 16, 16])
    rows = [
        ("Mașina electrică", "Puterea maximă (vârf)", "kW", F2, "Pmax_kW", "0.0"),
        ("", "Puterea nominală (Vmax constant)", "kW", F2, "Pnom_kW", "0.0"),
        ("", "Cuplul maxim la ax", "Nm", F2, "Max", "0.0"),
        ("", "Turația de bază", "rot/min", F2, "n_ax_Va", "0"),
        ("", "Turația maximă", "rot/min", None, "=c_n_max", "0"),
        ("Reductor", "Raportul i_rd = Ωr/Ωax", "-", F2, "i_rd", "0.00000"),
        ("", "Raportul de reducere n_motor/n_roată", "-", F2, "k_rd", "0.000"),
        ("Vehicul", "Accelerația realizabilă (aderență)", "m/s²", F2, "a", "0.000"),
        ("", "Timpul 0–100 km/h realizabil", "s", F2, "ta", "0.00"),
        ("Sistem de stocare", "Energia necesară în baterie", "kWh", F3, "Wbat_kWh", "0.00"),
        ("", "Consumul echivalent al ciclului P3", "kWh/100 km", F3, "cons", "0.00"),
        ("", "Celule (serie × paralel = total)", "-", None, "cfg", "@"),
        ("", "Energia instalată", "kWh", FB, "Wreal", "0.00"),
        ("", "Tensiunea nominală (min – max)", "V", None, "U", "@"),
        ("Convertor / invertor", "Curentul maxim al bateriei", "A", FB, "Ibat", "0"),
        ("", "Masa celulelor / pachetului (estimare)", "kg", None, "m", "@"),
    ]
    r = 5
    for sub, et, um, F, cheie, fmt in rows:
        ws.cell(row=r, column=1, value=sub).font = F_BOLD
        ws.cell(row=r, column=2, value=et).font = F_BASE
        ws.cell(row=r, column=3, value=um).font = F_BASE
        for k, L in enumerate(("D", "E")):
            if F is not None:
                val = f"={F.ref(cheie, k)}"
            elif cheie == "cfg":
                val = (f'=TEXT({FB.ref("Ns", k)},"0")&"s × "&TEXT({FB.ref("Np", k)},"0")&"p = "'
                       f'&TEXT({FB.ref("N", k)},"0")')
            elif cheie == "U":
                val = (f'=TEXT({FB.ref("Unom", k)},"0")&" ("&TEXT({FB.ref("Umin", k)},"0")&" – "'
                       f'&TEXT({FB.ref("Umax", k)},"0")&")"')
            elif cheie == "m":
                val = f'=TEXT({FB.ref("mcel", k)},"0")&" / "&TEXT({FB.ref("mpac", k)},"0")'
            else:
                val = cheie
            c = ws[f"{L}{r}"]
            c.value = val
            c.font = F_BOLD
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
        for cc in range(1, 6):
            ws.cell(row=r, column=cc).border = BORDER
        r += 1


def foaie_surse(wb, surse):
    ws = wb.create_sheet("Surse")
    titlu(ws, "Surse pentru datele de intrare (accesate 07–14.10.2026)")
    antet(ws, 4, ["Element", "Sursă", "URL"], [34, 60, 110])
    for i, (el, s, url) in enumerate(surse, 5):
        ws.cell(row=i, column=1, value=el).font = F_BASE
        ws.cell(row=i, column=2, value=s).font = F_BASE
        u = ws.cell(row=i, column=3, value=url)
        if url.startswith("http"):
            u.hyperlink = url
            u.font = F_LINK
        else:
            u.font = F_BASE
        for cc in range(1, 4):
            ws.cell(row=i, column=cc).border = BORDER


def construieste(tab, rez, C, cel, out, surse=()):
    wb = Workbook()
    col, rm, rgt = foaie_date(wb, tab)
    foaie_ipoteze(wb, C, cel)
    F2 = foaie_p2(wb, col, rm, rgt)
    F3 = foaie_p3(wb, col, rm, rgt, F2)
    for n in C["n_pasi"]:
        foaie_discretizare(wb, n, F3)
    FB = foaie_baterie(wb, F2, F3, col, rm, rgt)
    foaie_sinteza(wb, F2, F3, FB)
    foaie_surse(wb, surse)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    for F in (F2, F3, FB):
        F.ws.freeze_panes = "E5"
    wb.save(out)
    return dict(p2=F2.rand, p3=F3.rand, baterie=FB.rand)
