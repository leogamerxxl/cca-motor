"""Etapa 1 de calcul – dinamica longitudinală a vehiculului (curs CCA C1).

Construiește CCA_Calcul1_Dinamica_GT63.xlsx pentru Mercedes-AMG GT 63 4MATIC+ 4-Door Coupé:
date de intrare, calculul regimurilor caracteristice după relațiile (1)–(50) din
„Noțiuni generale de dinamica vehiculului” (L. Popescu, UPB), caracteristica de tracțiune
F(v) / P(v) și sinteza rezultatelor.

Utilizare: python3 build_calcul.py CCA_Calcul1_Dinamica_GT63.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

OUT = sys.argv[1]

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

IN = "Date de intrare"
CALC = "Calcul regimuri"
CAR = "Caracteristica tractiune"
REZ = "Rezultate"

wb = Workbook()


def header(ws, row, labels, widths=None):
    for c, h in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
        cell.border = BORDER
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def name(nm, sheet, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{sheet}'!{ref}")


# Number-to-text for labels and conclusions, independent of the reader's Excel locale:
# integers grouped with a space (2 582), decimals with a comma (2,39).
def n0(x):
    return f'SUBSTITUTE(SUBSTITUTE(FIXED({x},0),","," "),"."," ")'


def nd(x, d):
    return f'SUBSTITUTE(FIXED({x},{d},TRUE),".",",")'


def pct(x):
    return f'FIXED(100*({x}),0,TRUE)&" %"'


# ====================================================================== Date de intrare
ws = wb.active
ws.title = IN
ws["A1"] = "Etapa 1 de calcul – dinamica longitudinală: date de intrare"
ws["A1"].font = F_TITLE
ws["A2"] = ("Vehicul: Mercedes-AMG GT 63 4MATIC+ 4-Door Coupé (versiunea de serie a Concept AMG GT XX). "
            "Text albastru = valoare introdusă (modificabilă); text negru = formulă. "
            "Numele din coloana A pot fi folosite direct în formule; de exemplu, în foaia „Calcul regimuri” "
            "F_ra = ½·ρ·Cx·S·v² este =0.5*rho_aer*C_x*A_f*E7^2.")
ws["A2"].font = F_NOTE
ws.merge_cells("A2:F2")
ws.row_dimensions[2].height = 28
ws["A2"].alignment = WRAP
header(ws, 4, ["Nume", "Mărime", "Valoare", "Unitate", "Tip", "Sursă / justificare"], [11, 46, 12, 10, 12, 95])

OFFICIAL = "oficial"
COURSE = "curs C1"
ASSUME = "ipoteză"
CALC_T = "calculat"
SRC_MB = "Mercedes-Benz Media, comunicat de lansare 20.05.2026 (tabel tehnic)"

# (name, label, value-or-formula, unit, type, source, number format)
inputs = [
    ("SEC", "Vehicul", None, None, None, None, None),
    ("M_DIN", "Masa proprie DIN (fără șofer)", 2460, "kg", OFFICIAL, SRC_MB, "#,##0"),
    ("m_sofer", "Masa șoferului (convenția UE)", 75, "kg", ASSUME, "Regulamentul de punere în aplicare (UE) 2021/535 (anterior 1230/2012): masa în ordine de mers include șoferul de 75 kg", "0"),
    ("M_calc", "Masa totală de calcul M", "=M_DIN+m_sofer", "kg", CALC_T, "Egală cu masa UE în ordine de mers: 2535 kg (valoare oficială)", "#,##0"),
    ("C_x", "Coeficient de rezistență aerodinamică Cx", 0.22, "–", OFFICIAL, SRC_MB, "0.00"),
    ("A_f", "Aria frontală S", 2.44, "m²", OFFICIAL, SRC_MB + " (valoare publicată pentru gama AMG GT 4-Door Coupé electric)", "0.00"),
    ("T_max", "Cuplul maxim al sistemului de propulsie T_m,max", 2000, "N·m", OFFICIAL, SRC_MB + "; considerat la arborii motoarelor (vezi ipoteza transmisiei)", "#,##0"),
    ("P_vf", "Puterea de vârf (AMG Launch Control)", 860, "kW", OFFICIAL, SRC_MB + "; disponibilă cel mult 63 s, la SoC ≥ 80 %", "#,##0"),
    ("P_cont", "Puterea continuă", 530, "kW", OFFICIAL, SRC_MB, "#,##0"),
    ("v_max", "Viteza maximă (limitată electronic)", 300, "km/h", OFFICIAL, SRC_MB + " (cu AMG Driver's Package)", "0"),
    ("t_acc", "Timp de accelerare 0–100 km/h", 2.4, "s", OFFICIAL, SRC_MB + " (2,1 s este cu „1-foot rollout”)", "0.0"),
    ("t_200", "Timp de accelerare 0–200 km/h", 6.8, "s", OFFICIAL, SRC_MB, "0.0"),
    ("SEC", "Roată și transmisie", None, None, None, None, None),
    ("D_jant", "Diametrul jantei", 21, "inch", ASSUME, "Jante de 19″–21″ (oficial); se alege 21″, varianta cea mai performantă", "0"),
    ("B_anv", "Lățimea anvelopei (punte spate)", 295, "mm", ASSUME, "Dimensiunile anvelopelor versiunii electrice nu sunt publicate; se preia 295/30 R21 de la AMG GT 63 4-Door Coupé (V8). De confirmat în CoC.", "0"),
    ("h_rap", "Raportul de aspect al anvelopei (înălțime/lățime)", 0.30, "–", ASSUME, "Idem (profil 30)", "0.00"),
    ("k_def", "Coeficient de deformare radială sub sarcină", 0.97, "–", ASSUME, "Raza dinamică ≈ 0,96–0,98 × raza liberă (valoare uzuală)", "0.00"),
    ("r_w", "Raza dinamică a roții r_w", "=k_def*(D_jant*25.4/2+B_anv*h_rap)/1000", "m", CALC_T, "Raza liberă = D_jant·25,4/2 + B_anv·h_rap [mm]; r_w = k_def × raza liberă / 1000 [m]", "0.000"),
    ("n_m_vmax", "Turația motoarelor de pe puntea spate la v_max", 13000, "rot/min", OFFICIAL, "Mercedes-Benz Media: „peste 13 000 rot/min” la v_max (spate), peste 15 000 rot/min (față)", "#,##0"),
    ("i_g", "Raportul total de transmitere i_g (punte spate, estimat)", "=n_m_vmax/(v_max/3.6/(2*PI()*r_w)*60)", "–", CALC_T, "Raportul nu este publicat: i_g = n_m / n_w la v_max (limită inferioară, turația oficială fiind „peste 13 000”). Se folosește ca raport echivalent pentru întreaga acționare. Deoarece i_g se calculează din r_w, raportul i_g/r_w este fix: forțele, puterile și timpii nu depind de r_w (doar T_w și n_w).", "0.00"),
    ("eta_t", "Randamentul transmisiei η (motor → roată)", 0.95, "–", ASSUME, "Reductor planetar cu o treaptă (~0,97) × pierderi în arbori/lagăre; valoare uzuală 0,93–0,97", "0.00"),
    ("delta_m", "Coeficientul maselor în rotație δ", 1.0, "–", COURSE, "Curs C1, rel. (23): pentru modelul inițial F_i = M·a este suficient (δ = 1)", "0.00"),
    ("SEC", "Mediu și drum", None, None, None, None, None),
    ("g_acc", "Accelerația gravitațională g", 9.81, "m/s²", "constantă", "–", "0.00"),
    ("rho_aer", "Densitatea aerului ρ", 1.225, "kg/m³", ASSUME, "Atmosfera standard: 15 °C, 101,325 kPa, fără vânt (rel. 7)", "0.000"),
    ("f_r", "Coeficientul de rezistență la rulare f", 0.012, "–", COURSE, "Curs C1, tabel: asfalt uscat, normal – 0,012 (interval 0,010–0,015)", "0.000"),
    ("mu_us", "Coeficientul de aderență μ – asfalt uscat", 0.85, "–", COURSE, "Curs C1, tabel: asfalt uscat 0,80–0,90 (mijlocul intervalului)", "0.00"),
    ("mu_perf", "Coeficientul de aderență μ – anvelope de performanță", 1.20, "–", "calibrare", "Valoare de calibrare: μ necesar pentru 0–100 km/h în 2,4 s este 1,19 (regimul S2a). Depășește tabelul cursului (asfalt uscat 0,80–0,90) – realizabil doar cu anvelope sport pe asfalt uscat, cald.", "0.00"),
    ("P_cal", "Puterea medie disponibilă care reproduce 0–200 km/h", 750, "kW", "calibrare", "Aleasă astfel încât modelul (μ = 1,20) să dea 0–200 km/h ≈ 6,8 s – foaia „Caracteristica tractiune”, coloanele V–X. Arată că cei 860 kW nu sunt disponibili integral pe tot intervalul 100–200 km/h.", "#,##0"),
    ("SEC", "Cerințe de calcul (regimuri)", None, None, None, None, None),
    ("v_acc", "Viteza de capăt a accelerării", 100, "km/h", OFFICIAL, "0–100 km/h", "0"),
    ("a_med", "Accelerația medie impusă 0–100 km/h", "=v_acc/3.6/t_acc", "m/s²", CALC_T, "a = Δv / t", "0.00"),
    ("p_1", "Pantă de verificare (regim S3)", 10, "%", COURSE, "Curs C1, exemplul rel. (21): pantă de 10 %", "0"),
    ("v_p1", "Viteza constantă pe panta p_1", 100, "km/h", ASSUME, "Drum național / montan în rampă; panta de 10 % este exemplul din curs, rel. (21)", "0"),
    ("p_max", "Panta maximă pentru pornire (regim S4)", 30, "%", ASSUME, "Cerință uzuală pentru autoturisme (rampe de parcare); se precizează în tema de proiectare", "0"),
]

r = 5
for nm, label, val, unit, typ, src, fmt in inputs:
    if nm == "SEC":
        c = ws.cell(row=r, column=1, value=label)
        c.font = F_BOLD
        for cc in range(1, 7):
            ws.cell(row=r, column=cc).fill = FILL_SEC
            ws.cell(row=r, column=cc).border = BORDER
        r += 1
        continue
    ws.cell(row=r, column=1, value=nm).font = F_BOLD
    ws.cell(row=r, column=2, value=label).font = F_BASE
    v = ws.cell(row=r, column=3, value=val)
    v.font = F_BASE if (isinstance(val, str) and val.startswith("=")) else F_INPUT
    v.number_format = fmt
    v.alignment = RIGHT
    ws.cell(row=r, column=4, value=unit).font = F_BASE
    ws.cell(row=r, column=5, value=typ).font = F_BASE
    s = ws.cell(row=r, column=6, value=src)
    s.font = F_BASE
    s.alignment = WRAP
    for cc in range(1, 7):
        ws.cell(row=r, column=cc).border = BORDER
    name(nm, IN, f"$C${r}")
    r += 1

r += 1
ws.cell(row=r, column=1, value="Ipoteze ale modelului (curs C1):").font = F_BOLD
notes = [
    "Model longitudinal simplificat: vehiculul este un corp rigid care se deplasează pe direcția pantei; vânt nul (rel. 7, nu rel. 8).",
    "Tracțiune integrală (3 motoare: 2 spate + 1 față) cu distribuție ideală a cuplului între punți, proporțională cu sarcina dinamică a fiecărei punți (rel. 39); suma sarcinilor normale este F_z = M·g·cos α (rel. 37). Verificarea pe punți (înălțimea centrului de greutate, rapoartele față/spate – cap. 10) rămâne pentru etapa următoare.",
    "Acționarea este tratată ca un motor echivalent cu cuplul T_m,max și raportul i_g: F_t,drive = min(T_m,max·i_g·η/r_w ; η·P/v) (rel. 6 și 30), folosit în (40).",
    "f și μ sunt considerați constanți (nivelul „simplificat” din tabelul cursului: dimensionare preliminară).",
    "Masele în rotație (δ, rel. 23): la limita de aderență forța din contactul roată–drum accelerează doar masa M, a = (μ·F_z − F_rez)/M; cuplul suplimentar pentru masele în rotație îl furnizează motorul. Cu δ = 1 cele două forme coincid.",
]
for k, t in enumerate(notes, 1):
    c = ws.cell(row=r + k, column=1, value=f"{k}. {t}")
    c.font = F_BASE
ws.freeze_panes = "A5"

# ====================================================================== Calcul regimuri
cs = wb.create_sheet(CALC)
cs["A1"] = "Calculul regimurilor caracteristice – relațiile din cursul C1"
cs["A1"].font = F_TITLE
cs["A2"] = ("Fiecare coloană este un regim de funcționare. Rândurile urmează succesiunea din curs: "
            "rezistențe → forța de tracțiune necesară → cuplu și putere → verificarea aderenței și a acționării.")
cs["A2"].font = F_NOTE
cs.merge_cells("A2:I2")

scen = [
    ("S1", f'="Viteză maximă "&{n0("v_max")}&" km/h, drum orizontal"', "=v_max", 0, 0),
    ("S2a", f'="Accelerare 0–"&{n0("v_acc")}&" km/h în "&{nd("t_acc", 1)}&" s – la pornire (v = 0)"', 0, 0, "=a_med"),
    ("S2b", f'="Accelerare 0–"&{n0("v_acc")}&" km/h în "&{nd("t_acc", 1)}&" s – la "&{n0("v_acc")}&" km/h"', "=v_acc", 0, "=a_med"),
    ("S3", f'="Rampă "&{n0("p_1")}&" % la "&{n0("v_p1")}&" km/h, viteză constantă"', "=v_p1", "=p_1", 0),
    ("S4", f'="Pornire pe panta maximă de "&{n0("p_max")}&" % (v ≈ 0, a = 0)"', 0, "=p_max", 0),
]
HR = 4
cols = ["Rel. curs", "Mărime", "Expresie", "Unitate"] + [s[0] for s in scen]
header(cs, HR, cols, [9, 40, 34, 10] + [15] * len(scen))
cs.cell(row=HR + 1, column=2, value="Regim").font = F_BOLD
for j, s in enumerate(scen):
    c = cs.cell(row=HR + 1, column=5 + j, value=s[1])
    c.font = F_BOLD
    c.alignment = CENTER
    c.fill = FILL_KEY
    c.border = BORDER
cs.row_dimensions[HR + 1].height = 54

# (eq, label, expression text, unit, formula template using {c} for column letter and {r_*} row refs, fmt, key)
R = {}
rows = [
    ("", "Viteza v", "intrare", "km/h", "v_kmh", "0", False),
    ("", "Viteza v", "v / 3,6", "m/s", "={c}{v_kmh}/3.6", "0.00", False),
    ("", "Panta p = tg α", "intrare", "%", "p", "0", False),
    ("(18)", "Unghiul pantei α", "arctg(p/100)", "°", "=DEGREES(ATAN({c}{p}/100))", "0.00", False),
    ("", "Accelerația a", "intrare", "m/s²", "a", "0.00", False),
    ("(7)", "Rezistența aerodinamică F_ra", "½·ρ·Cx·S·v²", "N", "=0.5*rho_aer*C_x*A_f*{c}{v_ms}^2", "#,##0", False),
    ("(14)", "Rezistența la rulare F_f", "f·M·g·cos α", "N", "=f_r*M_calc*g_acc*COS(RADIANS({c}{alpha}))", "#,##0", False),
    ("(17)", "Rezistența la pantă G_t", "M·g·sin α", "N", "=M_calc*g_acc*SIN(RADIANS({c}{alpha}))", "#,##0", False),
    ("(22)/(23)", "Forța de inerție F_i", "δ·M·a", "N", "=delta_m*M_calc*{c}{a}", "#,##0", False),
    ("(25)/(26)", "Forța de tracțiune necesară F_t", "F_ra + F_f + G_t + M·a", "N", "=SUM({c}{fra}:{c}{fi})", "#,##0", True),
    ("(31)/(46)", "Cuplul la roți T_w", "F_t·r_w", "N·m", "={c}{ft}*r_w", "#,##0", True),
    ("(28)/(47)", "Puterea la roți P_w", "F_t·v", "kW", "={c}{ft}*{c}{v_ms}/1000", "#,##0.0", True),
    ("(29)/(30)", "Puterea cerută motoarelor P_m", "P_w / η", "kW", "={c}{pw}/eta_t", "#,##0.0", True),
    ("(33)", "Cuplul cerut motoarelor T_m", "F_t·r_w / (i_g·η)", "N·m", "={c}{ft}*r_w/(i_g*eta_t)", "#,##0", True),
    ("", "Turația roții n_w", "v / (2π·r_w) · 60", "rot/min", "={c}{v_ms}/(2*PI()*r_w)*60", "#,##0", False),
    ("", "Turația motoarelor n_m", "n_w · i_g", "rot/min", "={c}{nw}*i_g", "#,##0", True),
    ("(37)", "Sarcina normală pe roți F_z", "M·g·cos α", "N", "=M_calc*g_acc*COS(RADIANS({c}{alpha}))", "#,##0", False),
    ("(35)", "Forța maximă transmisibilă prin aderență μ·F_z", '="μ·F_z (μ = "&' + nd("mu_us", 2) + '&")"', "N", "=mu_us*{c}{fz}", "#,##0", False),
    ("(36)", "Coeficient de aderență necesar μ_nec", "F_t / F_z", "–", "={c}{ft}/{c}{fz}", "0.00", True),
    ("(36)/(41)", "Verificare aderență F_t ≤ μ·F_z", "", "", '=IF({c}{ft}<={c}{adh},"DA","NU – patinare")', "@", True),
    ("(6)/(30) → (40)", "Forța maximă a acționării F_t,drive", "min(T_m,max·i_g·η/r_w ; η·P_vârf/v)", "N",
     "=IF({c}{v_ms}=0,T_max*i_g*eta_t/r_w,MIN(T_max*i_g*eta_t/r_w,eta_t*P_vf*1000/{c}{v_ms}))", "#,##0", False),
    ("(42)", "Forța de tracțiune maximă realizabilă F_t,max", "min(F_t,drive ; μ·F_z)", "N", "=MIN({c}{fdrive},{c}{adh})", "#,##0", False),
    ("", "Limita activă pentru F_t,max", "", "", '=IF({c}{ftmax}>={c}{ft},"– (rezervă)",IF({c}{fdrive}<={c}{adh},"acționare","aderență"))', "@", False),
    ("(43)", "Condiția de realizare F_t,max ≥ F_t", "", "", '=IF({c}{ftmax}>={c}{ft},"ÎNDEPLINITĂ","NEÎNDEPLINITĂ")', "@", True),
    ("", "Rezerva de forță F_t,max − F_t", "", "N", "={c}{ftmax}-{c}{ft}", "#,##0;[Red]-#,##0", False),
    ("", "Încărcarea față de puterea continuă", "P_m / P_cont", "%", "={c}{pm}/P_cont", "0%", False),
    ("", "Încărcarea față de puterea de vârf", "P_m / P_vârf", "%", "={c}{pm}/P_vf", "0%", False),
]
keys = ["v_kmh", "v_ms", "p", "alpha", "a", "fra", "ff", "gt", "fi", "ft", "tw", "pw", "pm", "tm", "nw", "nm",
        "fz", "adh", "munec", "chk_adh", "fdrive", "ftmax", "limit", "cond", "rez", "load_c", "load_p"]
first = HR + 2
for k, key in enumerate(keys):
    R[key] = first + k

for k, (eq, label, expr, unit, tmpl, fmt, key_row) in enumerate(rows):
    row = first + k
    cs.cell(row=row, column=1, value=eq).font = F_BASE
    cs.cell(row=row, column=2, value=label).font = F_BOLD if key_row else F_BASE
    cs.cell(row=row, column=3, value=expr).font = F_BASE
    cs.cell(row=row, column=4, value=unit).font = F_BASE
    for j, s in enumerate(scen):
        col = 5 + j
        L = get_column_letter(col)
        if tmpl in ("v_kmh", "p", "a"):
            val = {"v_kmh": s[2], "p": s[3], "a": s[4]}[tmpl]
            cell = cs.cell(row=row, column=col, value=val)
            cell.font = F_BASE if isinstance(val, str) else F_INPUT
        else:
            cell = cs.cell(row=row, column=col, value=tmpl.format(c=L, **R))
            cell.font = F_BOLD if key_row else F_BASE
        cell.number_format = fmt
        cell.alignment = RIGHT if fmt != "@" else CENTER
        cell.border = BORDER
        if key_row:
            cell.fill = FILL_KEY
    for cc in range(1, 5):
        cs.cell(row=row, column=cc).border = BORDER
cs.freeze_panes = cs.cell(row=first, column=5)

# Derived indicators (closed-form)
d0 = first + len(rows) + 2
cs.cell(row=d0, column=1, value="Indicatori derivați").font = F_SUB
derived = [
    ("D1", '="Accelerația maximă la pornire, limitată de aderență (μ = "&' + nd("mu_us", 2) + '&")"',
     "(μ − f)·g  (forța de aderență accelerează doar masa M)", "m/s²", "=(mu_us-f_r)*g_acc", "0.00"),
    ("D2", '="Panta maximă urcată, limitată de aderență (μ = "&' + nd("mu_us", 2) + '&"), v ≈ 0"', "tg α = μ − f  (din μ·cos α ≥ f·cos α + sin α)", "%",
     "=(mu_us-f_r)*100", "0.0"),
    ("D3", "Raportul F_t,drive(0) / (M·g)", "T_m,max·i_g·η / (r_w·M·g)", "–",
     "=T_max*i_g*eta_t/(r_w*M_calc*g_acc)", "0.00"),
    ("D4", "Panta maximă limitată de acționare", "f·cos α + sin α = D3 (nelimitată dacă D3 > √(1+f²))", "%",
     '=IF(E{d3}>SQRT(1+f_r^2),"nelimitată",100*TAN(ASIN(E{d3}/SQRT(1+f_r^2))-ATAN(f_r)))', "0.0"),
    ("D5", "Viteza maximă limitată de puterea continuă (teoretic)", "½ρCxS·v³ + fMg·v = η·P_cont  (Cardano)", "km/h",
     "=3.6*({cub_cont})", "0"),
    ("D6", "Viteza maximă limitată de puterea de vârf (teoretic)", "½ρCxS·v³ + fMg·v = η·P_vârf  (Cardano)", "km/h",
     "=3.6*({cub_vf})", "0"),
    ("D7", "Viteza la care se trece de la limitarea de cuplu la cea de putere (P_vârf)", "v_b = η·P_vârf / F_t,drive(0)", "km/h",
     "=3.6*eta_t*P_vf*1000/(T_max*i_g*eta_t/r_w)", "0.0"),
    ("D8", '="Viteza la care forța limitată de puterea de vârf egalează limita de aderență (μ = "&' + nd("mu_perf", 2) + '&")"', "v = η·P_vârf / (μ_perf·M·g)", "km/h",
     "=3.6*eta_t*P_vf*1000/(mu_perf*M_calc*g_acc)", "0.0"),
]
header_row = d0 + 1
for c, h in enumerate(["Nr.", "Mărime", "Expresie", "Unitate", "Valoare"], 1):
    cell = cs.cell(row=header_row, column=c, value=h)
    cell.font = F_HEAD
    cell.fill = FILL_HEAD
    cell.alignment = CENTER
    cell.border = BORDER


def cardano(power_name):
    # Depressed cubic v^3 + p v - q = 0, with p = fMg/k, q = ηP/k, k = ½ρCxS -> single real root
    k = "(0.5*rho_aer*C_x*A_f)"
    p = f"(f_r*M_calc*g_acc/{k})"
    q = f"(eta_t*{power_name}*1000/{k})"
    disc = f"SQRT(({q})^2/4+({p})^3/27)"
    return f"POWER({q}/2+{disc},1/3)-POWER({disc}-{q}/2,1/3)"


DROW = {}
for k, (nr, label, expr, unit, f, fmt) in enumerate(derived):
    DROW[nr] = header_row + 1 + k
for k, (nr, label, expr, unit, f, fmt) in enumerate(derived):
    row = header_row + 1 + k
    f = f.format(d3=DROW["D3"], cub_cont=cardano("P_cont"), cub_vf=cardano("P_vf"))
    vals = [nr, label, expr, unit, f]
    for c, v in enumerate(vals, 1):
        cell = cs.cell(row=row, column=c, value=v)
        cell.font = F_BASE
        cell.border = BORDER
        cell.alignment = WRAP if c in (2, 3) else RIGHT
    cs.cell(row=row, column=5).number_format = fmt
    cs.cell(row=row, column=5).fill = FILL_KEY
    cs.row_dimensions[row].height = 26

# ====================================================================== Caracteristica de tracțiune
ts = wb.create_sheet(CAR)
ts["A1"] = "Caracteristica de tracțiune F(v), puterea P(v) și estimarea accelerării"
ts["A1"].font = F_TITLE
ts["A2"] = ("Pas de 5 km/h. F_t,drive = min(T_m,max·i_g·η/r_w ; η·P/v). F_rez = F_ra + F_f. "
            "a = min[(μ·M·g − F_rez)/M ; (F_t,drive − F_rez)/(δ·M)]; timpul se obține prin integrare numerică: "
            "Δt = 2·Δv/(a_i + a_i+1) (regula trapezelor aplicată ecuației dv/dt = a). "
            "Curba pentru puterea continuă folosește T_m,max la viteze mici, deoarece cuplul continuu nu este publicat. "
            "Coloanele V–X: sensibilitate cu puterea medie P_cal în locul puterii de vârf.")
ts["A2"].font = F_NOTE
ts.merge_cells("A2:X2")
ts.row_dimensions[2].height = 30
ts["A2"].alignment = WRAP
MU1, MU2 = nd("mu_us", 2), nd("mu_perf", 2)
cols_t = ["v [km/h]", "v [m/s]", "F_ra [N]", "F_f [N]", "F_rez orizontal [N]", f'="F_rez rampă "&{n0("p_1")}&" % [N]"',
          "F_t,drive vârf [N]", "F_t,drive continuu [N]", f'="μ·M·g (μ = "&{MU1}&") [N]"', f'="μ·M·g (μ = "&{MU2}&") [N]"',
          f'="F_t,max μ = "&{MU1}&" [N]"', f'="F_t,max μ = "&{MU2}&" [N]"', f'="a μ = "&{MU1}&" [m/s²]"',
          f'="a μ = "&{MU2}&" [m/s²]"', f'="t μ = "&{MU1}&" [s]"', f'="t μ = "&{MU2}&" [s]"', "P_m orizontal [kW]",
          f'="P_m rampă "&{n0("p_1")}&" % [kW]"', "P_vârf [kW]", "P_cont [kW]", "n_m [rot/min]",
          f'="F_t,drive P_cal = "&{n0("P_cal")}&" kW [N]"', f'="a μ = "&{MU2}&", P_cal [m/s²]"', f'="t μ = "&{MU2}&", P_cal [s]"']
TH = 4
header(ts, TH, cols_t, [8, 8, 9, 8, 11, 11, 11, 11, 11, 11, 11, 11, 9, 9, 9, 9, 11, 11, 9, 9, 10, 12, 10, 10])
ts.row_dimensions[TH].height = 42
STEP = 5
VMAXT = 300
t0 = TH + 1
n_pts = VMAXT // STEP + 1
for k in range(n_pts):
    row = t0 + k
    f = {
        1: STEP * k,
        2: f"=A{row}/3.6",
        3: f"=0.5*rho_aer*C_x*A_f*B{row}^2",
        4: "=f_r*M_calc*g_acc",
        5: f"=C{row}+D{row}",
        6: f"=C{row}+M_calc*g_acc*(f_r*COS(ATAN(p_1/100))+SIN(ATAN(p_1/100)))",
        7: f"=IF(B{row}=0,T_max*i_g*eta_t/r_w,MIN(T_max*i_g*eta_t/r_w,eta_t*P_vf*1000/B{row}))",
        8: f"=IF(B{row}=0,T_max*i_g*eta_t/r_w,MIN(T_max*i_g*eta_t/r_w,eta_t*P_cont*1000/B{row}))",
        9: "=mu_us*M_calc*g_acc",
        10: "=mu_perf*M_calc*g_acc",
        11: f"=MIN(G{row},I{row})",
        12: f"=MIN(G{row},J{row})",
        13: f"=MIN((I{row}-E{row})/M_calc,(G{row}-E{row})/(delta_m*M_calc))",
        14: f"=MIN((J{row}-E{row})/M_calc,(G{row}-E{row})/(delta_m*M_calc))",
        15: 0 if k == 0 else f"=O{row - 1}+(B{row}-B{row - 1})/((M{row}+M{row - 1})/2)",
        16: 0 if k == 0 else f"=P{row - 1}+(B{row}-B{row - 1})/((N{row}+N{row - 1})/2)",
        17: f"=E{row}*B{row}/eta_t/1000",
        18: f"=F{row}*B{row}/eta_t/1000",
        19: "=P_vf",
        20: "=P_cont",
        21: f"=B{row}/(2*PI()*r_w)*60*i_g",
        22: f"=IF(B{row}=0,T_max*i_g*eta_t/r_w,MIN(T_max*i_g*eta_t/r_w,eta_t*P_cal*1000/B{row}))",
        23: f"=MIN((J{row}-E{row})/M_calc,(V{row}-E{row})/(delta_m*M_calc))",
        24: 0 if k == 0 else f"=X{row - 1}+(B{row}-B{row - 1})/((W{row}+W{row - 1})/2)",
    }
    fmts = {1: "0", 2: "0.00", 13: "0.00", 14: "0.00", 15: "0.00", 16: "0.00", 23: "0.00", 24: "0.00"}
    for c in range(1, 25):
        cell = ts.cell(row=row, column=c, value=f[c])
        cell.font = F_INPUT if c == 1 else F_BASE
        cell.number_format = fmts.get(c, "#,##0")
        cell.border = BORDER
tl = t0 + n_pts - 1
ts.freeze_panes = ts.cell(row=t0, column=2)


def scatter(title, ytitle, series_cols, anchor, ymax=None):
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    ch.x_axis.title = "v [km/h]"
    ch.y_axis.title = ytitle
    ch.x_axis.scaling.min = 0
    ch.x_axis.scaling.max = VMAXT
    ch.y_axis.scaling.min = 0
    if ymax:
        ch.y_axis.scaling.max = ymax
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.height = 10
    ch.width = 18
    xref = Reference(ts, min_col=1, min_row=t0, max_row=tl)
    for col in series_cols:
        yref = Reference(ts, min_col=col, min_row=TH, max_row=tl)
        s = Series(yref, xref, title_from_data=True)
        s.marker.symbol = "none"
        s.smooth = False
        ch.series.append(s)
    ts.add_chart(ch, anchor)


scatter("Caracteristica de tracțiune: forțe la roți", "F [N]", [5, 6, 7, 8, 9, 10], f"Z{TH}", ymax=35000)
scatter("Puterea cerută motoarelor (v constantă) și puterea disponibilă", "P [kW]", [17, 18, 19, 20], f"Z{TH + 22}", ymax=900)

# ====================================================================== Rezultate
rs = wb.create_sheet(REZ, 0)
rs["A1"] = ("Etapa 1 de calcul – rezultate: dinamica longitudinală a autoturismului electric "
            "Mercedes-AMG GT 63 4MATIC+ 4-Door Coupé (AMG.EA)")
rs["A1"].font = F_TITLE
rs["A2"] = ("Calcul după „Noțiuni generale de dinamica vehiculului” (curs CCA C1, L. Popescu, UPB). "
            "Toate valorile sunt formule legate de foile „Date de intrare”, „Calcul regimuri” și „Caracteristica tractiune”.")
rs["A2"].font = F_NOTE
rs.merge_cells("A2:L2")

rs["A4"] = "1. Date principale"
rs["A4"].font = F_SUB
main = [
    ("Masa de calcul M", "=M_calc", "kg", "#,##0"),
    ("Cx · S", "=C_x*A_f", "m²", "0.000"),
    ("Coeficient de rezistență la rulare f / de aderență μ", "=" + nd("f_r", 3) + '&" / "&' + nd("mu_us", 2), "–", "@"),
    ("Raza dinamică a roții r_w", "=r_w", "m", "0.000"),
    ("Raport de transmitere estimat i_g", "=i_g", "–", "0.00"),
    ("Randament transmisie η", "=eta_t", "–", "0.00"),
    ("Cuplu maxim / putere de vârf / putere continuă", "=" + n0("T_max") + '&" N·m / "&' + n0("P_vf") + '&" kW / "&' + n0("P_cont") + '&" kW"', "", "@"),
]
for k, (lab, f, unit, fmt) in enumerate(main):
    row = 5 + k
    rs.cell(row=row, column=1, value=lab).font = F_BASE
    c = rs.cell(row=row, column=4, value=f)
    c.font = F_BASE
    c.number_format = fmt
    c.alignment = RIGHT
    rs.cell(row=row, column=5, value=unit).font = F_BASE

r0 = 5 + len(main) + 1
rs.cell(row=r0, column=1, value="2. Regimuri caracteristice (rel. 25–33, 35–43)").font = F_SUB
rh = r0 + 1
res_cols = ["Regim", "Descriere", "v [km/h]", "Pantă [%]", "a [m/s²]", "F_t nec. [N]", "T_w [N·m]", "P_w [kW]",
            "P_m [kW]", "T_m [N·m]", "n_m [rot/min]", "μ necesar", "Verificare (43)", "Limita activă"]
for c, h in enumerate(res_cols, 1):
    cell = rs.cell(row=rh, column=c, value=h)
    cell.font = F_HEAD
    cell.fill = FILL_HEAD
    cell.alignment = CENTER
    cell.border = BORDER
rs.row_dimensions[rh].height = 30
map_rows = [None, None, "v_kmh", "p", "a", "ft", "tw", "pw", "pm", "tm", "nm", "munec", "cond", "limit"]
fmt_res = [None, None, "0", "0", "0.00", "#,##0", "#,##0", "#,##0.0", "#,##0.0", "#,##0", "#,##0", "0.00", "@", "@"]
for j, s in enumerate(scen):
    row = rh + 1 + j
    L = get_column_letter(5 + j)
    rs.cell(row=row, column=1, value=s[0]).font = F_BOLD
    rs.cell(row=row, column=2, value=f"='{CALC}'!{L}{HR + 1}").font = F_BASE
    for c in range(3, len(res_cols) + 1):
        cell = rs.cell(row=row, column=c, value=f"='{CALC}'!{L}{R[map_rows[c - 1]]}")
        cell.font = F_BASE
        cell.number_format = fmt_res[c - 1]
        cell.alignment = RIGHT if fmt_res[c - 1] != "@" else CENTER
    for c in range(1, len(res_cols) + 1):
        rs.cell(row=row, column=c).border = BORDER
    rs.row_dimensions[row].height = 28
    rs.cell(row=row, column=1).alignment = Alignment(vertical="center")
    rs.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="center")

i0 = rh + len(scen) + 2
rs.cell(row=i0, column=1, value="3. Indicatori de performanță (model vs. date oficiale)").font = F_SUB
ih = i0 + 1
for c, h in enumerate(["Indicator", "", "", "Calculat", "Oficial", "Observație"] + [""] * 8, 1):
    cell = rs.cell(row=ih, column=c, value=h)
    cell.font = F_HEAD
    cell.fill = FILL_HEAD
    cell.alignment = CENTER
    cell.border = BORDER
rs.merge_cells(start_row=ih, start_column=1, end_row=ih, end_column=3)
rs.merge_cells(start_row=ih, start_column=6, end_row=ih, end_column=14)


def tlookup(col, v):
    return f"=INDEX('{CAR}'!{col}{t0}:{col}{tl},MATCH({v},'{CAR}'!A{t0}:A{tl},0))"


CS = f"'{CALC}'!"
# (key, label formula, calculated, official, number format, observation)
ind = [
    ("t100_us", '="Timp 0–100 km/h, μ = "&' + MU1 + '&" (asfalt uscat, tabelul cursului) [s]"', tlookup("O", "v_acc"),
     "=t_acc", "0.00", '="aderența limitează accelerarea"'),
    ("t100_perf", '="Timp 0–100 km/h, μ = "&' + MU2 + '&" (anvelope de performanță) [s]"', tlookup("P", "v_acc"),
     "=t_acc", "0.00", f'="calibrare: μ_perf ≈ μ necesar = "&{nd(CS + "F" + str(R["munec"]), 2)}&", deci nu este o verificare independentă"'),
    ("t200_perf", '="Timp 0–200 km/h, μ = "&' + MU2 + '&", P_vârf = "&' + n0("P_vf") + '&" kW [s]"', tlookup("P", "200"),
     "=t_200", "0.00", '="model optimist: puterea de vârf este considerată disponibilă integral până la 200 km/h"'),
    ("t100_cal", '="Timp 0–100 km/h, μ = "&' + MU2 + '&", P_cal = "&' + n0("P_cal") + '&" kW [s]"', tlookup("X", "v_acc"),
     "=t_acc", "0.00", '="practic neschimbat: până la ≈ 100 km/h accelerarea este limitată de aderență"'),
    ("t200_cal", '="Timp 0–200 km/h, μ = "&' + MU2 + '&", P_cal = "&' + n0("P_cal") + '&" kW [s]"', tlookup("X", "200"),
     "=t_200", "0.00", '="P_cal = puterea medie disponibilă care reproduce timpul oficial"'),
    ("a0", '="Accelerația maximă la pornire, μ = "&' + MU1 + '&" [m/s²]"', f"={CS}E{DROW['D1']}", "=a_med", "0.00",
     '="oficial = accelerația medie 0–100 km/h (din "&' + nd("t_acc", 1) + '&" s)"'),
    ("munec", '="Coeficient de aderență necesar pentru 0–100 km/h în "&' + nd("t_acc", 1) + '&" s"', f"={CS}F{R['munec']}",
     "", "0.00", '="regimul S2a"'),
    ("vteor", "Viteza maximă teoretică cu puterea continuă [km/h]", f"={CS}E{DROW['D5']}", "=v_max", "0",
     '="v_max = "&' + n0("v_max") + '&" km/h este limitată electronic (oficial)"'),
    ("pmax", '="Panta maximă, limită de aderență, μ = "&' + MU1 + '&" [%]"', f"={CS}E{DROW['D2']}", "", "0.0",
     '="tracțiune integrală, distribuție ideală a cuplului"'),
    ("pvmax", '="Puterea motoarelor la v_max = "&' + n0("v_max") + '&" km/h [kW]"', f"={CS}E{R['pm']}", "=P_cont", "0.0",
     '="comparată cu puterea continuă"'),
]
IND = {}
for k, (key, lab, fc, fo, fmt, obs) in enumerate(ind):
    row = ih + 1 + k
    IND[key] = f"D{row}"
    rs.cell(row=row, column=1, value=lab).font = F_BASE
    a = rs.cell(row=row, column=4, value=fc)
    a.font = F_BOLD
    a.number_format = fmt
    a.fill = FILL_KEY
    b = rs.cell(row=row, column=5, value=fo if fo else "–")
    b.font = F_BASE
    b.number_format = fmt
    rs.cell(row=row, column=6, value=obs).font = F_NOTE
    for c in range(1, 15):
        rs.cell(row=row, column=c).border = BORDER
    rs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    rs.merge_cells(start_row=row, start_column=6, end_row=row, end_column=14)

c0 = ih + len(ind) + 2
rs.cell(row=c0, column=1, value="4. Concluzii").font = F_SUB
S = f"'{CALC}'!"
def T(*parts):
    return "=" + "&".join(parts)


def q(text):
    return '"' + text + '"'


concl = [
    T(q("1. La v_max = "), n0("v_max"), q(" km/h pe drum orizontal sunt necesari F_t = "), n0(f"{S}E{R['ft']}"),
      q(" N la roți, adică P_m = "), n0(f"{S}E{R['pm']}"), q(" kW – doar "), pct(f"{S}E{R['load_c']}"),
      q(" din puterea continuă; "), pct(f"{S}E{R['fra']}/{S}E{R['ft']}"),
      q(" din rezistență este aerodinamică. Viteza maximă nu este limitată de putere (cu puterea continuă ar fi posibilă teoretic ≈ "),
      n0(f"{S}E{DROW['D5']}"), q(" km/h), ci electronic, conform datelor oficiale. Turația de "), n0("n_m_vmax"),
      q(" rot/min la v_max este o dată de intrare (folosită pentru estimarea i_g), nu un rezultat al calculului.")),
    T(q("2. Accelerarea 0–"), n0("v_acc"), q(" km/h în "), nd("t_acc", 1), q(" s cere o forță de "), n0(f"{S}F{R['ft']}"),
      q(" N la pornire, deci μ ≥ "), nd(f"{S}F{R['munec']}", 2), q(". Cu μ = "), nd("mu_us", 2),
      q(" (asfalt uscat, tabelul cursului) se pot transmite doar "), n0(f"{S}F{R['adh']}"), q(" N din cei "),
      n0(f"{S}F{R['fdrive']}"), q(" N pe care îi poate furniza acționarea, iar timpul minim estimat este "),
      nd(IND["t100_us"], 1), q(" s: accelerarea este limitată de aderență, nu de acționare.")),
    T(q("3. Timpul oficial de "), nd("t_acc", 1), q(" s se obține numai cu μ ≈ "), nd(f"{S}F{R['munec']}", 2),
      q(" (anvelope sport pe asfalt uscat, peste valorile din tabelul cursului); cu μ = "), nd("mu_perf", 2),
      q(" modelul dă "), nd(IND["t100_perf"], 2),
      q(" s – aceasta este o calibrare a lui μ, nu o verificare independentă. Raza r_w și raportul i_g nu influențează forțele, puterile și timpii, deoarece i_g este calculat din r_w; ele afectează doar T_w și n_w.")),
    T(q("4. Peste ≈ "), n0(f"{S}E{DROW['D8']}"), q(" km/h accelerarea este limitată de putere. Cu P_vârf = "), n0("P_vf"),
      q(" kW (η·P = "), n0("eta_t*P_vf"), q(" kW la roți) modelul dă 0–200 km/h în "), nd(IND["t200_perf"], 2),
      q(" s, cu "), pct(f"1-{IND['t200_perf']}/t_200"), q(" mai repede decât valoarea oficială de "), nd("t_200", 1),
      q(" s; timpul oficial se obține cu o putere medie disponibilă de ≈ "), n0("P_cal"), q(" kW ("),
      nd(IND["t200_cal"], 2),
      q(" s). Cauze posibile: puterea de vârf nu este disponibilă integral pe tot intervalul 100–200 km/h, randamentul real scade la turații mari, iar masele în rotație au fost neglijate (δ = 1, rel. 23).")),
    T(q("5. La "), n0("v_acc"), q(" km/h, accelerația medie impusă ("), nd("a_med", 2), q(" m/s²) ar cere P_m = "),
      n0(f"{S}G{R['pm']}"), q(" kW ("), pct(f"{S}G{R['load_p']}"),
      q(" din puterea de vârf): accelerația reală scade odată cu viteza, iar puterea de vârf este necesară doar pentru performanța maximă (Launch Control).")),
    T(q("6. Rampa de "), n0("p_1"), q(" % la "), n0("v_p1"), q(" km/h cere "), n0(f"{S}H{R['pm']}"),
      q(" kW, iar pornirea pe panta de "), n0("p_max"), q(" % cere "), n0(f"{S}I{R['ft']}"), q(" N (μ necesar "),
      nd(f"{S}I{R['munec']}", 2), q("): ambele condiții sunt îndeplinite cu rezervă mare.")),
    T(q("7. Date de confirmat în etapele următoare: dimensiunea anvelopelor (r_w), rapoartele reale de transmitere față/spate, distribuția cuplului între punți (cap. 10, rel. 39), cuplul continuu al motoarelor și curba reală a puterii de vârf în funcție de viteză.")),
]
for k, f in enumerate(concl):
    row = c0 + 1 + k
    cell = rs.cell(row=row, column=1, value=f)
    cell.font = F_BASE
    cell.alignment = WRAP
    rs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
    rs.row_dimensions[row].height = 44

for c, w in enumerate([7, 40, 9, 10, 9, 11, 11, 10, 10, 10, 11, 9, 15, 11], 1):
    rs.column_dimensions[get_column_letter(c)].width = w
rs.column_dimensions["A"].width = 9

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4

wb.active = 0
wb.save(OUT)
print("saved", OUT)
