"""Graficele etapei 1 de calcul, citite din CCA_Calcul1_Dinamica_GT63.xlsx (valori recalculate).

Utilizare: python3 grafice_calcul1.py CCA_Calcul1_Dinamica_GT63.xlsx
Produce fig_forte_tractiune.png, fig_puteri.png și fig_accelerare.png în același director.
"""
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

SRC = sys.argv[1]
OUT_DIR = os.path.dirname(os.path.abspath(SRC))

# Paleta categorială validată (ordinea fixă a sloturilor); textul folosește tonuri neutre.
C = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"]
INK, INK2, GRID, SURF = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": INK2, "axes.labelcolor": INK,
    "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "figure.facecolor": SURF,
    "axes.facecolor": SURF, "legend.frameon": False, "lines.linewidth": 2,
})
ro = FuncFormatter(lambda x, _: f"{x:,.0f}".replace(",", " "))          # 30 000
ro1 = FuncFormatter(lambda x, _: f"{x:.1f}".replace(".", ","))          # 2,4

wb = load_workbook(SRC, data_only=True)
ts = wb["Caracteristica tractiune"]
rows = [r for r in ts.iter_rows(min_row=5, values_only=True) if isinstance(r[0], (int, float))]
col = lambda i: [r[i] for r in rows]
v = col(0)
F_rez0, F_rez10, F_vf, F_cont, F_us, F_perf = col(4), col(5), col(6), col(7), col(8), col(9)
t_us, t_perf = col(14), col(15)
P0, P10, Pvf, Pc = col(16), col(17), col(18), col(19)
names = {dn: wb.defined_names[dn] for dn in wb.defined_names}


def nv(n):
    sheet, ref = names[n].attr_text.split("!")
    return wb[sheet.strip("'")][ref.replace("$", "")].value


t_acc, t_200 = nv("t_acc"), nv("t_200")


def end_label(ax, x, y, text, color, dy=0, ha="left"):
    ax.annotate(text, (x, y), xytext=(4, dy), textcoords="offset points", color=INK, fontsize=8.5,
                va="center", ha=ha)
    ax.plot([x], [y], "o", ms=4, color=color)


# 1. Forțe la roți
fig, ax = plt.subplots(figsize=(8.4, 5.0))
ax.plot(v, F_vf, color=C[2], label="F$_{t,drive}$ – putere de vârf 860 kW")
ax.plot(v, F_cont, color=C[3], label="F$_{t,drive}$ – putere continuă 530 kW")
ax.plot(v, F_perf, color=C[5], ls="--", lw=1.6, label="μ·F$_z$, μ = 1,20 (anvelope de performanță)")
ax.plot(v, F_us, color=C[4], ls="--", lw=1.6, label="μ·F$_z$, μ = 0,85 (asfalt uscat)")
ax.plot(v, F_rez10, color=C[1], label="F$_{rez}$ – rampă 10 %")
ax.plot(v, F_rez0, color=C[0], label="F$_{rez}$ – drum orizontal")
ax.set_xlim(0, 300)
ax.set_ylim(0, 35000)
ax.yaxis.set_major_formatter(ro)
ax.set_xlabel("v [km/h]")
ax.set_ylabel("F [N]")
ax.set_title("Caracteristica de tracțiune – Mercedes-AMG GT 63 4-Door Coupé", loc="left", color=INK, fontsize=11)
end_label(ax, 300, F_rez0[-1], "orizontal", C[0], dy=-7)
end_label(ax, 300, F_rez10[-1], "rampă 10 %", C[1], dy=5)
ax.legend(loc="upper right", fontsize=8.5)
fig.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "fig_forte_tractiune.png"), dpi=150)

# 2. Puteri
fig, ax = plt.subplots(figsize=(8.4, 5.0))
ax.plot(v, Pvf, color=C[2], ls="--", lw=1.6, label="putere de vârf 860 kW")
ax.plot(v, Pc, color=C[3], ls="--", lw=1.6, label="putere continuă 530 kW")
ax.plot(v, P10, color=C[1], label="P$_m$ – rampă 10 %")
ax.plot(v, P0, color=C[0], label="P$_m$ – drum orizontal")
ax.set_xlim(0, 300)
ax.set_ylim(0, 900)
ax.set_xlabel("v [km/h]")
ax.set_ylabel("P$_m$ [kW]")
ax.set_title("Puterea cerută motoarelor la viteză constantă", loc="left", color=INK, fontsize=11)
ax.plot([300], [P0[-1]], "o", ms=4, color=C[0])
ax.annotate(f"{P0[-1]:.0f} kW la 300 km/h", (300, P0[-1]), xytext=(-8, 14), textcoords="offset points",
            fontsize=8.5, color=INK, ha="right")
ax.plot([300], [P10[-1]], "o", ms=4, color=C[1])
ax.annotate(f"{P10[-1]:.0f} kW", (300, P10[-1]), xytext=(-8, 10), textcoords="offset points", fontsize=8.5,
            color=INK, ha="right")
ax.legend(loc="upper left", bbox_to_anchor=(0.0, 0.92), fontsize=8.5)
fig.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "fig_puteri.png"), dpi=150)

# 3. Accelerare
fig, ax = plt.subplots(figsize=(8.4, 5.0))
ax.plot(t_us, v, color=C[0], label="model, μ = 0,85 (asfalt uscat)")
ax.plot(t_perf, v, color=C[1], label="model, μ = 1,20 (anvelope de performanță)")
ax.plot([t_acc, t_200], [100, 200], "o", ms=8, color=INK, mfc=SURF, mew=2, label="date oficiale (0–100 / 0–200 km/h)")
i100, i200 = v.index(100), v.index(200)
for t, vv, color, dx, dy, ha in [(t_us[i100], 100, C[0], 8, -4, "left"), (t_perf[i100], 100, C[1], -8, 10, "right"),
                                 (t_perf[i200], 200, C[1], -8, 10, "right")]:
    ax.plot([t], [vv], "o", ms=6, color=color)
    ax.annotate(f"model {t:.2f} s".replace(".", ","), (t, vv), xytext=(dx, dy), textcoords="offset points",
                fontsize=8.5, color=INK, ha=ha)
for t, vv in [(t_acc, 100), (t_200, 200)]:
    ax.annotate(f"oficial {t:.1f} s".replace(".", ","), (t, vv), xytext=(8, -12), textcoords="offset points",
                fontsize=8.5, color=INK)
ax.set_xlim(0, 15)
ax.set_ylim(0, 310)
ax.xaxis.set_major_formatter(ro1)
ax.set_xlabel("t [s]")
ax.set_ylabel("v [km/h]")
ax.set_title("Accelerare din loc: model vs. date oficiale", loc="left", color=INK, fontsize=11)
ax.legend(loc="lower right", fontsize=8.5)
fig.tight_layout()
fig.savefig(os.path.join(OUT_DIR, "fig_accelerare.png"), dpi=150)
print("ok")
